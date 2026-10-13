"""
01_merge_raw.py
Menggabungkan 17 file mentah BPS "Month-to-Month Inflation (M-to-M)" (2010-2026)
menjadi satu file long-format tanpa transformasi (arsip mentah untuk transparansi).

Sumber: BPS, tabel "Inflasi Bulanan (M-to-M)"
        https://www.bps.go.id/en/statistics-table/2/MSMy/inflasi-bulanan-m-to-m-.html
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]

SCRIPT_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = SCRIPT_DIR / "raw"
OUTPUT_DIR = SCRIPT_DIR / "output"
YEARS = range(2010, 2027)

records = []
for yr in YEARS:
    path = UPLOAD_DIR / f"Month-to-Month Inflation (M-to-M), {yr}.xlsx"
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # baris data dimulai setelah 4 baris header
    for row in rows[4:]:
        if not row or row[0] is None:
            continue
        region = str(row[0]).strip()
        if not region:
            continue
        for i, month_name in enumerate(MONTHS):
            val = row[1 + i]
            if val is None or str(val).strip() in ("-", ""):
                continue
            records.append({
                "year": yr,
                "month": i + 1,
                "month_name": month_name,
                "region": region,
                "infl_pct_mtm": float(val),
            })

raw = pd.DataFrame(records).sort_values(["region", "year", "month"]).reset_index(drop=True)
OUTPUT_DIR.mkdir(exist_ok=True)
out_path = OUTPUT_DIR / "bps_ihk_raw_merged_2010_2026.csv"
raw.to_csv(out_path, index=False)

print(f"Total baris        : {len(raw)}")
print(f"Jumlah wilayah unik : {raw.region.nunique()}")
print(f"Rentang tahun       : {raw.year.min()}-{raw.year.max()}")
print(f"File output         : {out_path}")
print(f"Link                : {out_path.as_uri()}")
